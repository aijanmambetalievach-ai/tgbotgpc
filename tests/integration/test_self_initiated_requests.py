"""Requests a purchaser raises on their own, without a procurement."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.exceptions import PermissionDeniedError, ValidationError
from app.models.enums import PurchaseRequestStatus
from app.models.user import User
from app.services.excel_export_service import ExcelExportService
from app.services.purchase_request_service import PurchaseRequestService
from tests.integration.conftest import make_open_procurement
from tests.integration.test_purchase_request_flow import parse_items

pytestmark = pytest.mark.integration

SUBJECT = "Канцелярия для бухгалтерии на сентябрь"


async def create_and_submit(
    session: AsyncSession, purchaser: User, subject: str = SUBJECT, text: str = "Бумага 20 пачек"
):
    service = PurchaseRequestService(session)
    draft = await service.create_self_initiated_draft(purchaser, subject)
    await service.replace_items(purchaser, draft.request.id, parse_items(text))
    return await service.submit(purchaser, draft.request.id)


async def test_purchaser_creates_request_without_procurement(
    session: AsyncSession, purchasers: list[User]
) -> None:
    result = await create_and_submit(session, purchasers[0])

    assert result.procurement is None
    assert result.request.is_self_initiated is True
    assert result.request.procurement_id is None
    assert result.request.subject == SUBJECT
    assert result.request.status is PurchaseRequestStatus.SUBMITTED
    assert result.request.submitted_at is not None
    assert result.request.request_number == "REQ-2026-000001"
    assert len(result.items) == 1
    assert result.items[0].unit == "пач"


async def test_no_procurement_no_deadline_or_participation_checks(
    session: AsyncSession, purchasers: list[User]
) -> None:
    """A purchaser who is in no procurement at all can still raise a request."""
    lonely = purchasers[2]
    result = await create_and_submit(session, lonely, "Срочная потребность")
    assert result.request.status is PurchaseRequestStatus.SUBMITTED


async def test_repeated_start_reuses_the_open_draft(
    session: AsyncSession, purchasers: list[User]
) -> None:
    service = PurchaseRequestService(session)
    first = await service.create_self_initiated_draft(purchasers[0], SUBJECT)
    second = await service.create_self_initiated_draft(purchasers[0], "Другая тема")

    assert first.is_new is True
    assert second.is_new is False
    assert first.request.id == second.request.id
    assert second.request.subject == "Другая тема"


async def test_new_draft_allowed_after_previous_one_is_submitted(
    session: AsyncSession, purchasers: list[User]
) -> None:
    submitted = await create_and_submit(session, purchasers[0])
    service = PurchaseRequestService(session)

    second = await service.create_self_initiated_draft(purchasers[0], "Вторая тема")
    assert second.is_new is True
    assert second.request.id != submitted.request.id

    await service.replace_items(purchasers[0], second.request.id, parse_items("Соль 10 кг"))
    result = await service.submit(purchasers[0], second.request.id)
    assert result.request.request_number == "REQ-2026-000002"


async def test_subject_is_required(session: AsyncSession, purchasers: list[User]) -> None:
    service = PurchaseRequestService(session)
    for bad in ("", "  ", "ы"):
        with pytest.raises(ValidationError):
            await service.create_self_initiated_draft(purchasers[0], bad)


async def test_empty_self_initiated_request_cannot_be_submitted(
    session: AsyncSession, purchasers: list[User]
) -> None:
    service = PurchaseRequestService(session)
    draft = await service.create_self_initiated_draft(purchasers[0], SUBJECT)
    with pytest.raises(ValidationError):
        await service.submit(purchasers[0], draft.request.id)


async def test_double_submit_is_idempotent(
    session: AsyncSession, purchasers: list[User]
) -> None:
    service = PurchaseRequestService(session)
    draft = await service.create_self_initiated_draft(purchasers[0], SUBJECT)
    await service.replace_items(purchasers[0], draft.request.id, parse_items("Соль 10 кг"))

    first = await service.submit(purchasers[0], draft.request.id)
    second = await service.submit(purchasers[0], draft.request.id)

    assert first.already_submitted is False
    assert second.already_submitted is True
    assert first.request.request_number == second.request.request_number


async def test_employee_sees_it_among_new_requests(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    result = await create_and_submit(session, purchasers[0])
    page = await PurchaseRequestService(session).list_new(employee)

    assert result.request.id in {request.id for request in page.items}
    listed = next(r for r in page.items if r.id == result.request.id)
    assert listed.procurement is None
    assert listed.subject == SUBJECT


async def test_employee_approves_via_normal_statuses(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    submitted = await create_and_submit(session, purchasers[0])
    service = PurchaseRequestService(session)

    in_progress = await service.change_status(
        employee, submitted.request.id, PurchaseRequestStatus.IN_PROGRESS
    )
    assert in_progress.changed is True
    assert in_progress.procurement is None
    assert in_progress.purchaser.id == purchasers[0].id

    completed = await service.change_status(
        employee, submitted.request.id, PurchaseRequestStatus.COMPLETED
    )
    assert completed.new_status is PurchaseRequestStatus.COMPLETED

    history = await service.requests.list_history(submitted.request.id)
    assert [row.new_status for row in history] == [
        PurchaseRequestStatus.SUBMITTED,
        PurchaseRequestStatus.IN_PROGRESS,
        PurchaseRequestStatus.COMPLETED,
    ]


async def test_rejection_still_requires_a_reason(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    submitted = await create_and_submit(session, purchasers[0])
    service = PurchaseRequestService(session)

    with pytest.raises(ValidationError):
        await service.change_status(
            employee, submitted.request.id, PurchaseRequestStatus.REJECTED
        )

    rejected = await service.change_status(
        employee,
        submitted.request.id,
        PurchaseRequestStatus.REJECTED,
        comment="Нет бюджета в этом квартале.",
    )
    assert rejected.request.rejection_reason == "Нет бюджета в этом квартале."


async def test_purchaser_cannot_read_another_self_initiated_request(
    session: AsyncSession, purchasers: list[User]
) -> None:
    submitted = await create_and_submit(session, purchasers[0])
    service = PurchaseRequestService(session)

    with pytest.raises(PermissionDeniedError):
        await service.get_for_actor(purchasers[1], submitted.request.id)
    with pytest.raises(PermissionDeniedError):
        await ExcelExportService(session).export_request(
            purchasers[1], submitted.request.id
        )


async def test_excel_export_of_a_self_initiated_request(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    submitted = await create_and_submit(session, purchasers[0])
    excel = await ExcelExportService(session).export_request(employee, submitted.request.id)

    sheet = load_workbook(BytesIO(excel.content))["Заявка"]
    values = [
        cell.value
        for row in sheet.iter_rows(min_row=1)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert "Заявка по инициативе закупщика" in values
    assert SUBJECT in values
    assert "Бумага" in values


async def test_self_initiated_requests_stay_out_of_procurement_reports(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    await create_and_submit(session, purchasers[0], "Отдельная тема", "Гвозди 5 кг")

    stats = await PurchaseRequestService(session).procurements.stats(procurement.id)
    assert stats.responded == 0
    assert stats.total_items == 0

    aggregate = await ExcelExportService(session).export_procurement(employee, procurement.id)
    names = [
        cell.value
        for row in load_workbook(BytesIO(aggregate.content))["Заявки"].iter_rows(min_row=1)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert "Гвозди" not in names


async def test_both_kinds_coexist_for_one_purchaser(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    service = PurchaseRequestService(session)

    bound = await service.get_or_create_draft(purchasers[0], procurement.id)
    await service.replace_items(purchasers[0], bound.request.id, parse_items("Сахар 50 кг"))
    await service.submit(purchasers[0], bound.request.id)

    await create_and_submit(session, purchasers[0], "Своя тема", "Мука 10 кг")

    own = await service.list_own(purchasers[0])
    assert own.total == 2
    kinds = {request.is_self_initiated for request in own.items}
    assert kinds == {True, False}
