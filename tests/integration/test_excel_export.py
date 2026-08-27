"""Excel export: content, formatting and formula-injection safety."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.exceptions import PermissionDeniedError
from app.models.user import User
from app.services.excel_export_service import ExcelExportService
from tests.integration.conftest import make_open_procurement
from tests.integration.test_purchase_request_flow import submit_request

pytestmark = pytest.mark.integration


def read(content: bytes):
    return load_workbook(BytesIO(content))


async def test_individual_export(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    result = await submit_request(session, purchasers[0], procurement.id)

    excel = await ExcelExportService(session).export_request(employee, result.request.id)

    assert excel.filename.startswith(result.request.request_number or "")
    assert excel.filename.endswith(".xlsx")
    assert "Иван" in excel.filename  # Cyrillic preserved
    assert excel.size > 0

    workbook = read(excel.content)
    sheet = workbook["Заявка"]
    values = [
        [cell.value for cell in row] for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row)
    ]
    flat = [value for row in values for value in row if value is not None]

    assert procurement.procurement_number in flat
    assert "Сахар" in flat
    assert "кг" in flat
    assert sheet.freeze_panes is not None


async def test_quantities_are_numeric_cells(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    result = await submit_request(
        session, purchasers[0], procurement.id, "Масло 5,5 л\nСахар 50 кг"
    )

    workbook = read(
        (await ExcelExportService(session).export_request(employee, result.request.id)).content
    )
    sheet = workbook["Заявка"]

    quantities = []
    for row in sheet.iter_rows(min_row=1):
        for cell in row:
            if isinstance(cell.value, int | float | Decimal) and cell.column == 3:
                quantities.append(cell.value)
    assert any(float(value) == 5.5 for value in quantities)
    assert any(float(value) == 50 for value in quantities)


async def test_formula_injection_is_neutralised_in_excel(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    result = await submit_request(
        session,
        purchasers[0],
        procurement.id,
        "=SUM(A1:A5) 10 кг\n+cmd|calc 5 шт\n@formula 3 л",
    )

    workbook = read(
        (await ExcelExportService(session).export_request(employee, result.request.id)).content
    )
    sheet = workbook["Заявка"]
    for row in sheet.iter_rows(min_row=1):
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith(("=", "+", "@")), cell.value


async def test_aggregate_export_contains_every_purchaser(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    await submit_request(session, purchasers[0], procurement.id, "Сахар 50 кг\nМасло 60 л")
    await submit_request(session, purchasers[1], procurement.id, "Сахар 20 кг\nРис 80 кг")
    await submit_request(session, purchasers[2], procurement.id, "Яйца 500 шт")

    excel = await ExcelExportService(session).export_procurement(employee, procurement.id)
    workbook = read(excel.content)

    assert workbook.sheetnames == ["Заявки", "Сводка"]

    detail = workbook["Заявки"]
    names = [
        cell.value
        for row in detail.iter_rows(min_row=1)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert "Иван Иванов" in names
    assert "Азамат Садыков" in names
    assert "Данияр Алиев" in names
    assert "Рис" in names


async def test_summary_sheet_aggregates_identical_items(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    await submit_request(session, purchasers[0], procurement.id, "Сахар 50 кг\nМасло 60 л")
    await submit_request(session, purchasers[1], procurement.id, "сахар 20 кг\nРис 80 кг")
    await submit_request(session, purchasers[2], procurement.id, "Яйца 500 шт")

    excel = await ExcelExportService(session).export_procurement(employee, procurement.id)
    summary = read(excel.content)["Сводка"]

    totals: dict[str, float] = {}
    for row in summary.iter_rows(min_row=1):
        cells = [cell.value for cell in row]
        if len(cells) >= 3 and isinstance(cells[0], str) and isinstance(
            cells[2], int | float | Decimal
        ):
            totals[cells[0].lower()] = float(cells[2])

    # "Сахар" (50) + "сахар" (20) share a normalised name and unit -> merged.
    assert totals.get("сахар") == 70
    assert totals.get("масло") == 60
    assert totals.get("рис") == 80
    assert totals.get("яйца") == 500


async def test_different_names_are_never_merged(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    await submit_request(session, purchasers[0], procurement.id, "Масло 60 л")
    await submit_request(session, purchasers[1], procurement.id, "Масло подсолнечное 40 л")

    summary = read(
        (await ExcelExportService(session).export_procurement(employee, procurement.id)).content
    )["Сводка"]
    labels = [
        cell.value
        for row in summary.iter_rows(min_row=1)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert "Масло" in labels
    assert "Масло подсолнечное" in labels


async def test_same_name_different_unit_stays_separate(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    await submit_request(session, purchasers[0], procurement.id, "Масло 60 л")
    await submit_request(session, purchasers[1], procurement.id, "Масло 40 кг")

    summary = read(
        (await ExcelExportService(session).export_procurement(employee, procurement.id)).content
    )["Сводка"]
    rows = [
        (row[0].value, row[1].value, row[2].value)
        for row in summary.iter_rows(min_row=1)
        if row[0].value == "Масло"
    ]
    assert {(unit, float(total)) for _name, unit, total in rows} == {("л", 60.0), ("кг", 40.0)}


async def test_unicode_and_long_text_are_handled(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    result = await submit_request(
        session, purchasers[0], procurement.id, "Ёлка новогодняя «Люкс» 3 шт"
    )
    excel = await ExcelExportService(session).export_request(employee, result.request.id)
    values = [
        cell.value
        for row in read(excel.content)["Заявка"].iter_rows(min_row=1)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any("Ёлка новогодняя «Люкс»" in value for value in values)


async def test_rbac_on_exports(
    session: AsyncSession, employee: User, purchasers: list[User]
) -> None:
    procurement = await make_open_procurement(session, employee, purchasers)
    result = await submit_request(session, purchasers[0], procurement.id)
    service = ExcelExportService(session)

    # Aggregate export is employee-only.
    with pytest.raises(PermissionDeniedError):
        await service.export_procurement(purchasers[0], procurement.id)

    # A purchaser may export their own request…
    own = await service.export_request(purchasers[0], result.request.id)
    assert own.size > 0

    # …but never somebody else's.
    with pytest.raises(PermissionDeniedError):
        await service.export_request(purchasers[1], result.request.id)
