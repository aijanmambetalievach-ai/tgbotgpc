"""End-to-end acceptance scenario from the specification.

Employee creates a procurement (both ways) → picks purchasers → publishes →
each purchaser submits a list → employee sees requests, downloads individual
and aggregate Excel, changes statuses.
"""

from __future__ import annotations

import datetime as dt
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.timeutils import utcnow
from app.models.enums import ParticipantSelectionMode, ProcurementStatus, PurchaseRequestStatus
from app.models.user import User
from app.parsers.procurement_parser import RegexProcurementParser
from app.parsers.purchase_list_parser import RegexPurchaseListParser
from app.services.excel_export_service import ExcelExportService
from app.services.procurement_service import ProcurementDraftInput, ProcurementService
from app.services.purchase_request_service import PurchaseRequestService
from app.services.statistics_service import StatisticsService

pytestmark = pytest.mark.integration


async def test_full_scenario(
    session: AsyncSession, employee: User, purchasers: list[User], settings
) -> None:
    ivan, azamat, daniyar = purchasers
    procurements = ProcurementService(session)
    requests = PurchaseRequestService(session)

    # 1. Employee creates the procurement step by step and publishes it.
    published = await procurements.create_and_publish(
        employee,
        ProcurementDraftInput(
            title="Продукты питания на сентябрь",
            deadline=utcnow() + dt.timedelta(days=3),
            description="Сбор потребности подразделений.",
            participant_mode=ParticipantSelectionMode.MANUAL,
            participant_ids=[ivan.id, azamat.id, daniyar.id],
        ),
    )
    procurement = published.procurement
    assert procurement.status is ProcurementStatus.OPEN
    assert procurement.procurement_number == "TENDER-2026-000001"
    assert len(published.participants) == 3

    # 2. Each purchaser sees exactly this procurement and submits their list.
    lists = {
        ivan: "1. Сахар 50 кг\n2. Масло 60 литров\n3. Мука 100 кг",
        azamat: "1. Рис 80 кг\n2. Молоко 100 литров",
        daniyar: "1. Яйца 500 шт\n2. Картофель 150 кг",
    }
    parser = RegexPurchaseListParser()
    submitted_numbers = []
    for purchaser, text in lists.items():
        visible = await procurements.list_for_purchaser(purchaser)
        assert [item.id for item in visible.items] == [procurement.id]

        draft = await requests.get_or_create_draft(purchaser, procurement.id)
        parsed = parser.parse(text)
        assert parsed.errors == []
        await requests.replace_items(purchaser, draft.request.id, parsed.items)
        result = await requests.submit(purchaser, draft.request.id)
        submitted_numbers.append(result.request.request_number)

    assert submitted_numbers == ["REQ-2026-000001", "REQ-2026-000002", "REQ-2026-000003"]

    # 3. Employee dashboard numbers match the specification example.
    card = await procurements.get_card(employee, procurement.id)
    assert card.stats.participants == 3
    assert card.stats.responded == 3
    assert card.stats.total_items == 7
    assert card.stats.response_rate == 100

    # 4. Individual Excel.
    first_request = (await requests.list_by_procurement(employee, procurement.id)).items[0]
    single = await ExcelExportService(session).export_request(employee, first_request.id)
    assert single.filename.endswith(".xlsx")
    assert load_workbook(BytesIO(single.content)).sheetnames == ["Заявка"]

    # 5. Aggregate Excel with a summary sheet.
    aggregate = await ExcelExportService(session).export_procurement(employee, procurement.id)
    workbook = load_workbook(BytesIO(aggregate.content))
    assert workbook.sheetnames == ["Заявки", "Сводка"]
    detail_rows = sum(
        1
        for row in workbook["Заявки"].iter_rows(min_row=1)
        if isinstance(row[0].value, int)
    )
    assert detail_rows == 7

    # 6. Status change reaches the purchaser's request.
    change = await requests.change_status(
        employee, first_request.id, PurchaseRequestStatus.IN_PROGRESS
    )
    assert change.changed is True
    assert change.new_status is PurchaseRequestStatus.IN_PROGRESS
    assert change.purchaser.id in {ivan.id, azamat.id, daniyar.id}

    # 7. Global statistics.
    stats = await StatisticsService(session).dashboard(employee)
    assert stats.open_procurements == 1
    assert stats.active_purchasers == 3
    assert stats.requests_today == 3
    assert stats.in_progress_requests == 1
    assert stats.new_requests == 2


async def test_quick_mode_scenario(
    session: AsyncSession, employee: User, purchasers: list[User], settings
) -> None:
    parsed = RegexProcurementParser(
        timezone=settings.timezone, default_deadline_time=settings.default_deadline_time_obj
    ).parse(
        "Название: Канцелярия на сентябрь\n"
        f"Срок: {(utcnow() + dt.timedelta(days=6)).strftime('%d.%m.%Y')} 17:00\n"
        "Описание: Собрать потребность всех подразделений"
    )
    assert parsed.is_complete
    assert parsed.title == "Канцелярия на сентябрь"
    assert parsed.deadline is not None

    result = await ProcurementService(session).create_and_publish(
        employee,
        ProcurementDraftInput(
            title=parsed.title or "",
            deadline=parsed.deadline,
            description=parsed.description,
            participant_mode=ParticipantSelectionMode.ALL_ACTIVE,
        ),
    )

    assert result.procurement.status is ProcurementStatus.OPEN
    assert result.procurement.participant_mode is ParticipantSelectionMode.ALL_ACTIVE
    assert len(result.participants) == len(purchasers)
