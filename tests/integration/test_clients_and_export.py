"""Registration, bootstrap, client administration and Excel exports."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.settings import Settings
from app.domain.exceptions import PermissionDeniedError, ValidationError
from app.models.enums import ClientCategory, ConsultationTopic, UserRole
from app.models.user import User
from app.services.application_service import ApplicationService
from app.services.auth_service import AuthService
from app.services.client_service import ClientService
from app.services.excel_export_service import ExcelExportService
from app.services.statistics_service import StatisticsService

pytestmark = pytest.mark.integration


async def test_first_contact_registers_a_client(session: AsyncSession) -> None:
    auth = AuthService(session)
    user = await auth.resolve_or_register(5001, "aigul", "Айгуль")
    assert user.role is UserRole.CLIENT
    assert user.full_name == "Айгуль"
    assert user.category is None

    again = await auth.resolve_or_register(5001, "aigul_new", "Другое имя")
    assert again.id == user.id
    assert again.telegram_username == "aigul_new"
    assert again.full_name == "Айгуль", "the name a client gave us is not overwritten"


async def test_bootstrap_promotes_admins(session: AsyncSession, settings: Settings) -> None:
    existing = await AuthService(session).resolve_or_register(6002, None, "Будущий админ")
    configured = settings.model_copy(update={"bootstrap_admin_telegram_ids": "6001, 6002"})
    auth = AuthService(session, configured)

    changed = await auth.bootstrap_admins()
    assert {user.telegram_id for user in changed} == {6001, 6002}
    assert existing.role is UserRole.ADMIN
    assert await auth.bootstrap_admins() == [], "bootstrap is idempotent"

    fresh = await auth.resolve_or_register(6001, "boss", "Настоящее Имя")
    assert fresh.full_name == "Настоящее Имя"


async def test_profile_validation(session: AsyncSession, client: User) -> None:
    service = ClientService(session)
    await service.update_contacts(client, full_name="Айгуль С.", phone="0700 11 22 33")
    assert client.phone == "+996700112233"
    with pytest.raises(ValidationError):
        await service.update_contacts(client, phone="не помню")
    with pytest.raises(ValidationError):
        await service.update_contacts(client, full_name="А")


async def test_blocking(session: AsyncSession, admin: User, client: User) -> None:
    service = ClientService(session)
    await service.set_blocked(admin, client.id, True)
    assert not client.can_use_bot
    assert await service.broadcast_recipients(admin, None) == []
    await service.set_blocked(admin, client.id, False)
    assert [u.id for u in await service.broadcast_recipients(admin, ClientCategory.SUPPLIER)] == [
        client.id
    ]
    with pytest.raises(PermissionDeniedError):
        await service.set_blocked(admin, admin.id, True)
    with pytest.raises(PermissionDeniedError):
        await service.set_blocked(client, admin.id, True)


async def test_exports_and_statistics(session: AsyncSession, admin: User, client: User) -> None:
    await ApplicationService(session).request_consultation(
        client, ConsultationTopic.CONTRACT, '=HYPERLINK("http://evil") как расторгнуть договор?'
    )
    excel = await ExcelExportService(session).export_applications(admin)
    assert excel.filename.startswith("Заявки_") and excel.filename.endswith(".xlsx")
    sheet = load_workbook(BytesIO(excel.content)).active
    assert sheet is not None
    assert sheet.cell(row=2, column=5).value == "Договор и его исполнение"
    assert str(sheet.cell(row=2, column=13).value).startswith("'=HYPERLINK")

    clients = await ExcelExportService(session).export_clients(admin)
    client_sheet = load_workbook(BytesIO(clients.content)).active
    assert client_sheet is not None
    assert client_sheet.cell(row=2, column=1).value == client.full_name

    with pytest.raises(PermissionDeniedError):
        await ExcelExportService(session).export_clients(client)

    stats = await StatisticsService(session).dashboard(admin)
    assert stats.clients_total == 1
    assert stats.clients_by_category[ClientCategory.SUPPLIER] == 1
