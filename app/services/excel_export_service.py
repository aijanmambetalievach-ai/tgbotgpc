"""Excel exports for administrators (openpyxl).

Security: every string written to a cell goes through
:func:`app.domain.text.sanitize_excel_cell`, which neutralises formula
injection (``=``, ``+``, ``-``, ``@``, tab, CR) — client-typed text ends up in
these files.
"""

from __future__ import annotations

import logging
from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.settings import Settings, get_settings
from app.domain.permissions import require_admin
from app.domain.text import sanitize_excel_cell, sanitize_filename
from app.domain.timeutils import format_datetime, to_local, utcnow
from app.models.audit_log import AuditAction
from app.models.enums import (
    APPLICATION_STATUS_LABELS,
    APPLICATION_TYPE_LABELS,
    CATEGORY_LABELS,
    CONTACT_METHOD_LABELS,
    TOPIC_LABELS,
)
from app.models.user import User
from app.repositories.applications import ApplicationRepository
from app.repositories.audit import AuditRepository
from app.repositories.users import UserRepository

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="DDE7F0")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="9AA5B1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 8

APPLICATION_HEADERS = (
    "№ заявки",
    "Дата",
    "Тип",
    "Статус",
    "Курс / тема",
    "Клиент",
    "Категория",
    "Телефон",
    "Организация",
    "Telegram",
    "Способ связи",
    "Удобное время",
    "Описание",
    "Причина отклонения",
)
CLIENT_HEADERS = (
    "Имя",
    "Категория",
    "Телефон",
    "Организация",
    "Telegram",
    "Дата регистрации",
    "Доступ",
)


@dataclass(frozen=True, slots=True)
class ExcelFile:
    filename: str
    content: bytes

    @property
    def size(self) -> int:
        return len(self.content)


class ExcelExportService:
    def __init__(self, session: AsyncSession, settings: Settings | None = None) -> None:
        self.session = session
        self.settings = settings or get_settings()
        self.applications = ApplicationRepository(session)
        self.users = UserRepository(session)
        self.audit = AuditRepository(session)

    async def export_applications(self, actor: User) -> ExcelFile:
        require_admin(actor)
        tz = self.settings.timezone
        rows = await self.applications.list_for_export()
        workbook, sheet = self._new_sheet("Заявки", APPLICATION_HEADERS)
        for index, row in enumerate(rows, start=2):
            application, client = row.application, row.client
            subject = application.course_title or (
                TOPIC_LABELS[application.topic] if application.topic else ""
            )
            self._write_row(
                sheet,
                index,
                [
                    application.number,
                    format_datetime(application.created_at, tz),
                    APPLICATION_TYPE_LABELS[application.type],
                    APPLICATION_STATUS_LABELS[application.status],
                    subject,
                    client.full_name,
                    CATEGORY_LABELS[client.category] if client.category else "",
                    client.phone or "",
                    client.organization or "",
                    f"@{client.telegram_username}" if client.telegram_username else "",
                    # Drop the leading emoji — it is noise in a spreadsheet.
                    CONTACT_METHOD_LABELS[application.contact_method].split(" ", 1)[-1],
                    application.preferred_time or "",
                    application.description,
                    application.rejection_reason or "",
                ],
            )
        return await self._finish(
            actor, workbook, sheet, "Заявки", AuditAction.EXPORT_APPLICATIONS, len(rows)
        )

    async def export_clients(self, actor: User) -> ExcelFile:
        require_admin(actor)
        tz = self.settings.timezone
        clients = await self.users.list_all_clients()
        workbook, sheet = self._new_sheet("Клиенты", CLIENT_HEADERS)
        for index, client in enumerate(clients, start=2):
            self._write_row(
                sheet,
                index,
                [
                    client.full_name,
                    CATEGORY_LABELS[client.category] if client.category else "",
                    client.phone or "",
                    client.organization or "",
                    f"@{client.telegram_username}" if client.telegram_username else "",
                    format_datetime(client.created_at, tz),
                    "активен" if client.is_active else "заблокирован",
                ],
            )
        return await self._finish(
            actor, workbook, sheet, "Клиенты", AuditAction.EXPORT_CLIENTS, len(clients)
        )

    # -- helpers -----------------------------------------------------------
    @staticmethod
    def _new_sheet(title: str, headers: Sequence[str]) -> tuple[Workbook, Worksheet]:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = title
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        return workbook, sheet

    @staticmethod
    def _write_row(sheet: Worksheet, row: int, values: Sequence[object]) -> None:
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=sanitize_excel_cell(value))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _autofit(sheet: Worksheet) -> None:
        for column_cells in sheet.columns:
            longest = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=0,
            )
            letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[letter].width = max(
                MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, longest + 2)
            )

    async def _finish(
        self,
        actor: User,
        workbook: Workbook,
        sheet: Worksheet,
        name: str,
        action: str,
        count: int,
    ) -> ExcelFile:
        self._autofit(sheet)
        buffer = BytesIO()
        workbook.save(buffer)
        stamp = to_local(utcnow(), self.settings.timezone).strftime("%Y-%m-%d")
        filename = sanitize_filename(f"{name}_{stamp}") + ".xlsx"
        self.audit.log(
            action=action,
            entity_type="export",
            actor_user_id=actor.id,
            payload={"rows": count},
        )
        await self.session.commit()
        logger.info(
            "excel exported",
            extra={"operation": action, "user_id": actor.id, "result": f"rows={count}"},
        )
        return ExcelFile(filename=filename, content=buffer.getvalue())
