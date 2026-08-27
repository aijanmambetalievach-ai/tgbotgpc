"""Parser for purchase lists uploaded as an Excel workbook.

Accepts both shapes seen in practice:

* an **exported request** — a few metadata rows, then a header row
  (``№ | Наименование | Количество | Ед. измерения | Комментарий``) and data;
* a **bare table** — data starting on the first row, with or without headers,
  in any column order.

Whatever cannot be read confidently becomes a :class:`ParseError` for the user
to fix — exactly like the text parser. Nothing is ever guessed.

Safety notes:

* the workbook is opened ``read_only`` and ``data_only``, so formulas are never
  evaluated — only their cached values are read, and a formula without one
  simply yields an unreadable row;
* row scanning is capped (``max_rows``) and the caller caps the file size, which
  together bound the cost of a hostile archive.
"""

from __future__ import annotations

import logging
from collections.abc import Iterable, Sequence
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.domain.text import collapse_whitespace, normalize_unicode
from app.domain.units import normalize_unit
from app.parsers.base import ParsedItem, ParseError, ParseResult
from app.parsers.purchase_list_parser import (
    DEFAULT_MAX_ITEMS,
    DEFAULT_MAX_NAME_LENGTH,
    RegexPurchaseListParser,
)

logger = logging.getLogger(__name__)

DEFAULT_MAX_ROWS = 5_000
#: How far down we look for a header row before assuming there is none.
HEADER_SEARCH_DEPTH = 30

REASON_UNREADABLE_FILE = "не удалось прочитать файл — ожидается книга Excel (.xlsx)"
REASON_EMPTY_FILE = "в файле нет данных"
REASON_TOO_MANY_ROWS = "в файле слишком много строк"

_NAME_ALIASES = (
    "наименование", "название", "товар", "товары", "номенклатура", "позиция",
    "продукт", "материал", "предмет", "item", "name", "product",
)
_QUANTITY_ALIASES = (
    "количество", "кол-во", "колво", "кол", "кол.", "объем", "qty", "quantity", "amount",
)
_UNIT_ALIASES = (
    "ед", "ед.", "ед изм", "ед. изм", "ед.изм", "ед измерения", "ед. измерения",
    "единица", "единицы", "единица измерения", "unit", "uom", "измерение",
)
_COMMENT_ALIASES = (
    "комментарий", "примечание", "заметка", "описание", "comment", "note", "remark",
)
_INDEX_ALIASES = ("№", "n", "no", "номер", "п/п", "№ п/п", "#", "nn")


def _cell_text(value: Any) -> str:
    """Normalise a cell into trimmed text (numbers included)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return collapse_whitespace(normalize_unicode(str(value)))


def _header_key(value: Any) -> str:
    return _cell_text(value).lower().replace("ё", "е").strip(" .:")


def _matches(key: str, aliases: Sequence[str]) -> bool:
    if not key:
        return False
    cleaned = {alias.replace("ё", "е").strip(" .:") for alias in aliases}
    return key in cleaned or any(key.startswith(alias) for alias in cleaned if len(alias) > 3)


class ColumnMap:
    """Which column holds what, once a header row has been recognised."""

    __slots__ = ("comment", "index", "name", "quantity", "unit")

    def __init__(self) -> None:
        self.name: int | None = None
        self.quantity: int | None = None
        self.unit: int | None = None
        self.comment: int | None = None
        self.index: int | None = None

    @property
    def is_usable(self) -> bool:
        return self.name is not None and self.quantity is not None

    def get(self, row: Sequence[Any], column: int | None) -> Any:
        if column is None or column >= len(row):
            return None
        return row[column]


class ExcelPurchaseListParser:
    """Implementation of :class:`PurchaseListFileParserProtocol`."""

    def __init__(
        self,
        max_items: int = DEFAULT_MAX_ITEMS,
        max_name_length: int = DEFAULT_MAX_NAME_LENGTH,
        max_rows: int = DEFAULT_MAX_ROWS,
    ) -> None:
        self.max_items = max_items
        self.max_rows = max_rows
        self.text_parser = RegexPurchaseListParser(
            max_items=max_items, max_name_length=max_name_length
        )

    # -- public API --------------------------------------------------------
    def parse_file(self, content: bytes) -> ParseResult:
        try:
            rows = self._read_rows(content)
        except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
            logger.info(
                "rejected uploaded spreadsheet",
                extra={
                    "operation": "parse_excel_list",
                    "result": "unreadable",
                    "error_type": type(exc).__name__,
                },
            )
            return ParseResult(
                items=[], errors=[ParseError(line="", reason=REASON_UNREADABLE_FILE)]
            )

        if not rows:
            return ParseResult(items=[], errors=[ParseError(line="", reason=REASON_EMPTY_FILE)])

        header_index, columns = self._find_header(rows)
        data_rows = rows[header_index + 1 :] if header_index >= 0 else rows
        return self._parse_rows(data_rows, columns)

    # -- internals ---------------------------------------------------------
    def _read_rows(self, content: bytes) -> list[tuple[Any, ...]]:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        try:
            sheet = workbook.worksheets[0] if workbook.worksheets else None
            if sheet is None:
                return []
            rows: list[tuple[Any, ...]] = []
            for row in sheet.iter_rows(values_only=True):
                if len(rows) >= self.max_rows:
                    break
                rows.append(tuple(row))
            return rows
        finally:
            workbook.close()

    def _find_header(self, rows: Iterable[tuple[Any, ...]]) -> tuple[int, ColumnMap]:
        """Locate the header row. Returns ``(-1, empty_map)`` when absent."""
        for row_index, row in enumerate(rows):
            if row_index >= HEADER_SEARCH_DEPTH:
                break
            columns = ColumnMap()
            for column_index, value in enumerate(row):
                key = _header_key(value)
                if columns.name is None and _matches(key, _NAME_ALIASES):
                    columns.name = column_index
                elif columns.quantity is None and _matches(key, _QUANTITY_ALIASES):
                    columns.quantity = column_index
                elif columns.unit is None and _matches(key, _UNIT_ALIASES):
                    columns.unit = column_index
                elif columns.comment is None and _matches(key, _COMMENT_ALIASES):
                    columns.comment = column_index
                elif columns.index is None and key in _INDEX_ALIASES:
                    columns.index = column_index
            if columns.is_usable:
                return row_index, columns
        return -1, ColumnMap()

    def _parse_rows(
        self, rows: Sequence[tuple[Any, ...]], columns: ColumnMap
    ) -> ParseResult:
        items: list[ParsedItem] = []
        errors: list[ParseError] = []

        for row in rows:
            if len(items) >= self.max_items:
                errors.append(ParseError(line="", reason=REASON_TOO_MANY_ROWS))
                break
            if not any(_cell_text(value) for value in row):
                continue  # blank separator row

            parsed, error = (
                self._parse_mapped_row(row, columns)
                if columns.is_usable
                else self._parse_free_row(row)
            )
            if parsed is not None:
                items.append(parsed)
            elif error is not None:
                errors.append(error)

        if not items and not errors:
            errors.append(ParseError(line="", reason=REASON_EMPTY_FILE))
        return ParseResult(items=items, errors=errors)

    def _parse_mapped_row(
        self, row: Sequence[Any], columns: ColumnMap
    ) -> tuple[ParsedItem | None, ParseError | None]:
        name = _cell_text(columns.get(row, columns.name))
        quantity_raw = columns.get(row, columns.quantity)
        unit_text = _cell_text(columns.get(row, columns.unit))
        comment = _cell_text(columns.get(row, columns.comment)) or None

        source = " ".join(
            part for part in (name, _cell_text(quantity_raw), unit_text) if part
        )
        if not source:
            return None, None

        quantity = _to_decimal(quantity_raw)
        unit = normalize_unit(unit_text)

        if name and quantity is not None and quantity > 0 and unit:
            return (
                ParsedItem(
                    name=self._clip(name),
                    quantity=quantity,
                    unit=unit,
                    source_text=source,
                    confidence=1.0,
                    comment=comment,
                ),
                None,
            )

        # Anything less than perfectly typed goes through the text parser, which
        # already knows how to read "50 кг", "5,5 литров" and friends.
        parsed, error = self.text_parser.parse_line(source)
        if parsed is not None and comment:
            parsed = ParsedItem(
                name=parsed.name,
                quantity=parsed.quantity,
                unit=parsed.unit,
                source_text=parsed.source_text,
                confidence=parsed.confidence,
                comment=comment,
            )
        return parsed, error

    def _parse_free_row(
        self, row: Sequence[Any]
    ) -> tuple[ParsedItem | None, ParseError | None]:
        """No header: glue the cells together and reuse the text parser."""
        cells = [_cell_text(value) for value in row]
        meaningful = [cell for cell in cells if cell]
        if not meaningful:
            return None, None
        # A leading standalone integer is a row number, not a quantity.
        if len(meaningful) > 2 and meaningful[0].isdigit():
            meaningful = meaningful[1:]
        return self.text_parser.parse_line(" ".join(meaningful))

    def _clip(self, name: str) -> str:
        limit = self.text_parser.max_name_length
        return name if len(name) <= limit else name[:limit].rstrip()


def _to_decimal(value: Any) -> Decimal | None:
    """Exact decimal from a numeric cell; ``None`` when the cell is not a number."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # str() keeps the value the user typed (5.5 -> "5.5"), unlike Decimal(float).
        return Decimal(str(value))
    if isinstance(value, str):
        text = collapse_whitespace(value).replace(",", ".").replace(" ", "")
        try:
            return Decimal(text)
        except InvalidOperation:
            return None
    return None


def get_default_excel_parser(
    max_items: int = DEFAULT_MAX_ITEMS,
    max_name_length: int = DEFAULT_MAX_NAME_LENGTH,
    max_rows: int = DEFAULT_MAX_ROWS,
) -> ExcelPurchaseListParser:
    return ExcelPurchaseListParser(
        max_items=max_items, max_name_length=max_name_length, max_rows=max_rows
    )
